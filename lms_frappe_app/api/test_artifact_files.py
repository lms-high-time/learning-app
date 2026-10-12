# Copyright (c) 2026, NikoMusaev and Contributors
# See license.txt

import io

import frappe
from frappe.tests import IntegrationTestCase

from lms_frappe_app.agent_learning.artifacts import codes, write
from lms_frappe_app.agent_learning.artifacts.course import _схемы_курса
from lms_frappe_app.agent_learning.spaces import пространство_курса
from lms_frappe_app.api import manager, student
from lms_frappe_app.tests.sample_data import (
	добавить_в_организацию,
	зачислить,
	создать_менеджера,
	создать_организацию,
	создать_ученика,
	создать_урок,
	схема_документа,
)

CSV = "Месяц;Выручка;Расходы\nЯнварь;100;80\nФевраль;120;90\n".encode()


def xlsx(с_формулой: bool = True) -> bytes:
	from openpyxl import Workbook

	книга = Workbook()
	лист = книга.active
	лист.append(["Месяц", "Выручка", "Расходы", "Прибыль"])
	лист.append(["Январь", 100, 80, "=B2-C2" if с_формулой else 20])
	поток = io.BytesIO()
	книга.save(поток)
	return поток.getvalue()


class IntegrationTestArtifactFiles(IntegrationTestCase):
	"""Блок документа как файл или ссылка (learning-services#315)."""

	def setUp(self):
		self.addCleanup(frappe.set_user, "Administrator")
		суффикс = frappe.generate_hash(length=6)
		self.ученик = создать_ученика(f"file-{суффикс}@example.com")
		self.урок = создать_урок(f"Урок {суффикс}")
		self.курс = зачислить(self.ученик, self.урок)
		frappe.get_doc(
			{
				"doctype": "Agent Course Artifact",
				"course": self.курс,
				"slug": "plan",
				"title": "План",
				"blocks": [
					{"block_key": "money", "title": "Финплан", "kind": "file", "accept": "xlsx,csv", "lesson": self.урок},
					{"block_key": "crm", "title": "База клиентов", "kind": "link"},
					{"block_key": "goal", "title": "Цель"},
				],
			}
		).insert(ignore_permissions=True)
		frappe.set_user(self.ученик)

	def загрузить(self, имя: str = "plan.csv", данные: bytes = CSV, ключ: str = "money") -> dict:
		пространство = пространство_курса(self.ученик, self.курс)
		return student.контракт(write._положить_файл)(
			self.ученик, self.курс, пространство, "plan", ключ, имя, данные
		)

	def запись_файла(self, ключ: str = "money") -> str | None:
		"""Имя `File` в строке блока. По адресу искать нельзя: Frappe отдаёт
		одинаковому содержимому один `file_url` на несколько записей."""
		документ = frappe.db.get_value("Agent Student Artifact", {"student": self.ученик, "course": self.курс})
		return frappe.db.get_value(
			"Agent Artifact Content", {"parent": документ, "block_key": ключ}, "file"
		)

	def блок(self, ключ: str) -> dict:
		блоки = student.artifact(self.курс, "plan")["data"]["blocks"]
		return next(б for б in блоки if б["key"] == ключ)

	# --- файл ---

	def test_загруженный_csv_виден_файлом_и_срезом(self):
		"""Файл виден в блоке, перечне и блоках старта урока (`artifact_blocks`);
		документ заводится до файла: иначе Frappe снимает привязку (#343)."""
		from lms_frappe_app.agent_learning.artifacts.document import _блоки_урока

		ответ = self.загрузить()["data"]

		self.assertEqual(ответ["file"]["name"], "plan.csv")
		блок = self.блок("money")
		self.assertEqual(блок["kind"], "file")
		self.assertEqual(блок["accept"], ["xlsx", "csv"])
		self.assertEqual(блок["file"]["type"], "csv")
		self.assertIn("| Месяц | Выручка | Расходы |", блок["preview"])
		self.assertIn("| Февраль | 120 | 90 |", блок["preview"])
		перечень = student.artifact(self.курс)["data"]["artifacts"][0]
		self.assertEqual(перечень["blocks_filled"], 1, "файл без текста — блок заполнен")

		with self.subTest("#343 файл первой загрузки привязан к документу"):
			файл = frappe.get_doc("File", self.запись_файла())
			документ = frappe.db.get_value(
				"Agent Student Artifact", {"student": self.ученик, "course": self.курс}
			)
			self.assertEqual(
				(файл.attached_to_doctype, файл.attached_to_name), ("Agent Student Artifact", документ)
			)

		блоки = _блоки_урока(self.ученик, self.курс, self.урок, пространство_курса(self.ученик, self.курс))
		self.assertEqual([б["key"] for б in блоки], ["money"])
		self.assertEqual(блоки[0]["file"]["name"], "plan.csv")
		self.assertTrue(блоки[0]["preview"])

	def test_файл_base64_от_сервиса_ложится_так_же(self):
		import base64

		ответ = student.upload_artifact_file(
			self.курс, "plan", "money", file_name="plan.csv", content=base64.b64encode(CSV).decode()
		)

		self.assertTrue(ответ["ok"], ответ.get("error"))
		self.assertEqual(ответ["data"]["file"]["name"], "plan.csv")
		self.assertIn("| Месяц | Выручка | Расходы |", ответ["data"]["preview"])

	def test_отказы_записи_в_блок(self):
		"""Файл, таблица и ссылка — только в блок своего вида и только годные.

		Отказы файла — одной строкой проводки: правило целиком проверяет
		`agent_learning/artifacts/test_files.py`."""

		def записать(ключ: str, **поля) -> dict:
			return student.update_artifact(self.курс, "plan", ключ, **поля)

		отказы = {
			"файл не того типа": (lambda: self.загрузить("photo.png", b"\x89PNG"), codes.ФАЙЛ_НЕ_ТОГО_ТИПА),
			"таблица в текстовый блок": (lambda: записать("goal", table=[["a"]]), codes.ВИД_НЕ_ТОТ),
			"кривая таблица": (lambda: записать("money", table=["не строки"]), codes.НЕВЕРНАЯ_ТАБЛИЦА),
			"большая таблица": (lambda: записать("money", table=[["x"]] * 501), codes.НЕВЕРНАЯ_ТАБЛИЦА),
			"ссылка в текстовый блок": (
				lambda: записать("goal", url="https://example.com"),
				codes.ВИД_НЕ_ТОТ,
			),
			"кривая ссылка": (lambda: записать("crm", url="javascript:alert(1)"), codes.НЕВЕРНАЯ_ССЫЛКА),
		}
		for что, (вызов, код) in отказы.items():
			with self.subTest(что):
				ответ = вызов()
				self.assertEqual(ответ["error"]["code"], код)
				if что == "файл не того типа":
					self.assertEqual(ответ["error"]["accept"], ["xlsx", "csv"])

	def test_новый_файл_замещает_прежний(self):
		self.загрузить()
		первый = self.запись_файла()
		второй = self.загрузить("plan.xlsx", xlsx(с_формулой=False))["data"]["file"]

		self.assertEqual(self.блок("money")["file"]["name"], "plan.xlsx")
		self.assertFalse(frappe.db.exists("File", первый), "прежний файл удалён")
		self.assertTrue(второй["url"].startswith("/private/files/"), "файл приватный")

	def test_слишком_большой_файл_отказ(self):
		frappe.set_user("Administrator")
		frappe.db.set_single_value("Agent Learning Settings", "artifact_file_max_mb", 1)
		frappe.clear_document_cache("Agent Learning Settings", "Agent Learning Settings")
		self.addCleanup(frappe.clear_document_cache, "Agent Learning Settings", "Agent Learning Settings")
		self.addCleanup(frappe.db.set_single_value, "Agent Learning Settings", "artifact_file_max_mb", 10)
		frappe.set_user(self.ученик)

		ответ = self.загрузить(данные=b"a;b\n" * 400_000)

		self.assertEqual(ответ["error"]["code"], codes.ФАЙЛ_СЛИШКОМ_БОЛЬШОЙ)

	def test_очистка_удаляет_файл(self):
		self.загрузить()
		файл = self.запись_файла()

		student.update_artifact(self.курс, "plan", "money", clear=True)

		self.assertIsNone(self.блок("money")["file"])
		self.assertFalse(frappe.db.exists("File", файл))

	def test_файл_видит_только_ученик(self):
		self.загрузить()
		файл = frappe.get_doc("File", self.запись_файла())
		frappe.set_user("Administrator")
		суффикс = frappe.generate_hash(length=6)
		другой = создать_ученика(f"other-{суффикс}@example.com")
		организация = создать_организацию(f"Компания {суффикс}")
		добавить_в_организацию(self.ученик, организация)
		руководитель = создать_менеджера(f"mg-{суффикс}@example.com", организация)

		self.assertTrue(frappe.has_permission("File", doc=файл, user=self.ученик))
		for чужой in (другой, руководитель):
			with self.subTest(чужой=чужой):
				self.assertFalse(frappe.has_permission("File", doc=файл, user=чужой))

	def test_миграция_привязывает_потерянный_файл(self):
		from lms_frappe_app.patches.v0_1 import attach_artifact_files

		self.загрузить()
		имя = self.запись_файла()
		frappe.db.set_value("File", имя, {"attached_to_doctype": None, "attached_to_name": None})

		attach_artifact_files.execute()

		документ = frappe.db.get_value("Agent Student Artifact", {"student": self.ученик, "course": self.курс})
		self.assertEqual(frappe.db.get_value("File", имя, "attached_to_name"), документ)

	def test_файл_документа_организации_открывает_её_руководитель(self):
		"""Права на файл — от документа, документ курса от компании — её (#341, #343)."""
		frappe.set_user("Administrator")
		суффикс = frappe.generate_hash(length=6)
		организация = создать_организацию(f"Компания {суффикс}")
		добавить_в_организацию(self.ученик, организация)
		frappe.get_doc(
			{"doctype": "Course Allocation", "organization": организация, "course": self.курс}
		).insert(ignore_permissions=True)
		руководитель = создать_менеджера(f"mgf-{суффикс}@example.com", организация)
		frappe.set_user(self.ученик)
		self.загрузить()

		файл = frappe.get_doc("File", self.запись_файла())
		self.assertTrue(frappe.has_permission("File", doc=файл, user=руководитель))

	# --- таблица от агента (#319) ---

	def файл_блока(self, ключ: str = "money") -> bytes:
		return frappe.get_doc("File", self.запись_файла(ключ)).get_content()

	def test_таблица_агента_ложится_xlsx_с_формулой(self):
		from openpyxl import load_workbook

		ответ = student.update_artifact(
			self.курс,
			"plan",
			"money",
			table=[["Месяц", "Выручка", "Расходы", "Прибыль"], ["Январь", 100, 80, "=B2-C2"]],
			file_name="Финплан",
			content="Первый месяц в плюсе.",
		)

		self.assertTrue(ответ["ok"], ответ.get("error"))
		блок = self.блок("money")
		self.assertEqual(блок["file"]["name"], "Финплан.xlsx")
		self.assertEqual(блок["content"], "Первый месяц в плюсе.")
		self.assertIn("=B2-C2", блок["preview"])
		лист = load_workbook(io.BytesIO(self.файл_блока())).active
		self.assertEqual(лист["D2"].value, "=B2-C2", "формула осталась формулой")

	def test_блок_только_csv_получает_csv(self):
		frappe.set_user("Administrator")
		frappe.get_doc(
			{
				"doctype": "Agent Course Artifact",
				"course": self.курс,
				"slug": "log",
				"title": "Журнал",
				"blocks": [{"block_key": "calls", "title": "Звонки", "kind": "file", "accept": "csv"}],
			}
		).insert(ignore_permissions=True)
		frappe.set_user(self.ученик)

		ответ = student.update_artifact(self.курс, "log", "calls", table=[["Дата", "Клиент"], ["1.10", "Анна"]])

		self.assertTrue(ответ["ok"], ответ.get("error"))
		блоки = student.artifact(self.курс, "log")["data"]["blocks"]
		self.assertEqual(блоки[0]["file"]["name"], "calls.csv")
		self.assertIn("| Дата | Клиент |", блоки[0]["preview"])

	def test_таблица_строкой_json_тоже_принимается(self):
		ответ = student.update_artifact(self.курс, "plan", "money", table='[["Месяц"], ["Январь"]]')

		self.assertTrue(ответ["ok"], ответ.get("error"))
		self.assertEqual(self.блок("money")["file"]["name"], "money.xlsx")

		# --- ссылка ---

	def test_ссылка_записывается_в_блок_ссылку(self):
		ответ = student.update_artifact(self.курс, "plan", "crm", url="https://crm.example.com/base")

		self.assertTrue(ответ["ok"])
		self.assertEqual(self.блок("crm")["url"], "https://crm.example.com/base")
		self.assertEqual(ответ["data"]["blocks_filled"], 1)

	def test_текстовый_блок_как_прежде(self):
		student.update_artifact(self.курс, "plan", "goal", "Открыть вторую кофейню")

		блок = self.блок("goal")
		self.assertEqual((блок["kind"], блок["content"], блок["file"], блок["url"]), ("text", "Открыть вторую кофейню", None, None))

	# --- отчёт руководителя ---

	def test_отчёт_руководителя_считает_файл_заполненным(self):
		# Курс от организации — до первой записи: документ ложится в её
		# пространство, а личный отчёт организации не считает (#341).
		организация = создать_организацию(f"Отчёт {frappe.generate_hash(length=6)}")
		добавить_в_организацию(self.ученик, организация)
		frappe.get_doc(
			{"doctype": "Course Allocation", "organization": организация, "course": self.курс}
		).insert(ignore_permissions=True)
		self.загрузить()
		frappe.set_user("Administrator")
		руководитель = создать_менеджера(f"rep-{frappe.generate_hash(length=6)}@example.com", организация)
		frappe.set_user(руководитель)

		строка = next(с for с in manager.org_report()["data"]["rows"] if с["user"] == self.ученик)

		self.assertEqual(строка["document"], {"blocks_total": 3, "blocks_filled": 1})


class IntegrationTestArtifactKinds(IntegrationTestCase):
	"""Вид блока и форматы файла в схеме документа курса."""

	def setUp(self):
		суффикс = frappe.generate_hash(length=6)
		self.урок = создать_урок(f"Урок {суффикс}")
		self.курс = frappe.db.get_value(
			"Course Chapter", frappe.db.get_value("Course Lesson", self.урок, "chapter"), "course"
		)

	def схема(self, блоки: list[dict]) -> dict:
		return схема_документа(self.курс, "plan", "План", блоки)

	def test_вид_и_форматы_сохраняются(self):
		self.схема([{"key": "money", "title": "Финплан", "kind": "file", "accept": ".XLSX, csv"}])

		блок = _схемы_курса(self.курс)[0].blocks[0]

		self.assertEqual((блок.kind, блок.accept), ("file", "xlsx,csv"))

	def test_неизвестный_вид_отказ(self):
		ответ = self.схема([{"key": "money", "title": "Финплан", "kind": "pdf"}])

		self.assertEqual(ответ["error"]["code"], codes.НЕВЕРНЫЙ_ВИД_БЛОКА)
