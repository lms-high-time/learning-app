# Copyright (c) 2026, NikoMusaev and Contributors
# See license.txt

"""Документ курса через методы ученика (learning-services#515).

Логику разделов проверяет `agent_learning/artifacts/test_tables.py` без базы;
здесь — что она дошла до методов: запись строк и полей, чтение документа,
перечень и прогресс, очистка, отказы самой записи, версия схемы и
выгрузка.
"""

from datetime import datetime
from io import BytesIO
from unittest.mock import patch

import frappe
from frappe.tests import IntegrationTestCase

from lms_frappe_app.agent_learning.artifacts import codes, write
from lms_frappe_app.agent_learning.errors import Отказ
from lms_frappe_app.api import student
from lms_frappe_app.tests.sample_data import (
	зачислить,
	раздел,
	создать_урок,
	создать_ученика,
	схема_документа,
)

ЖУРНАЛ = раздел(
	"agreements",
	"Договорённости",
	"many",
	[
		{"key": "topic", "title": "Тема", "required": True},
		{"key": "decision", "title": "Решение", "required": False},
		{"key": "owner", "title": "Кто отвечает", "required": {"if_column": "decision"}},
	],
	description="Что команда решила",
)
ПРАВИЛА = раздел(
	"how_we_decide", "Как мы решаем", "one", [{"key": "method", "title": "Способ", "required": True}]
)


class IntegrationTestArtifactAPI(IntegrationTestCase):
	def setUp(self):
		self.addCleanup(frappe.set_user, "Administrator")
		суффикс = frappe.generate_hash(length=6)
		self.ученик = создать_ученика(f"doc-{суффикс}@example.com")
		self.урок = создать_урок(f"Урок {суффикс}")
		self.курс = зачислить(self.ученик, self.урок)
		self.схема([{**ЖУРНАЛ, "lesson": self.урок}, ПРАВИЛА])
		frappe.set_user(self.ученик)

	def схема(self, разделы: list, ключ: str = "agreements_log") -> str:
		frappe.set_user("Administrator")
		ответ = схема_документа(self.курс, ключ, "Журнал договорённостей", разделы, purpose="Зачем журнал")
		self.assertTrue(ответ["ok"], ответ)
		frappe.set_user(self.ученик)
		return ответ["data"]["id"]

	def записать(self, ключ: str, **параметры) -> dict:
		return student.update_artifact(self.курс, "agreements_log", ключ, **параметры)

	def документ(self) -> dict:
		return student.artifact(self.курс, "agreements_log")["data"]

	def test_запись_и_чтение(self):
		ответ = self.записать(
			"agreements", rows=[{"topic": "Ретро", "decision": "Раз в две недели"}, {"topic": "Созвоны"}]
		)
		self.assertTrue(ответ["ok"], ответ)
		self.assertEqual(
			ответ["data"],
			{
				"artifact": "agreements_log",
				"key": "agreements",
				"blocks_total": 2,
				"blocks_filled": 0,
				"empty_cells": [{"row": "A1", "column": "owner"}],
				"created": ["A1", "A2"],
			},
		)
		ответ = self.записать("Agreements", rows=[{"id": "A1", "owner": "Маша"}])
		self.assertEqual((ответ["data"]["blocks_filled"], ответ["data"]["created"]), (1, []))
		self.записать("how_we_decide", fields={"method": "Согласие"})

		документ = self.документ()
		self.assertEqual(
			{к: документ[к] for к in ("course", "artifact", "title", "purpose", "space")},
			{
				"course": self.курс,
				"artifact": "agreements_log",
				"title": "Журнал договорённостей",
				"purpose": "Зачем журнал",
				"space": "personal",
			},
		)
		журнал, правила = документ["blocks"]
		self.assertEqual(
			журнал,
			{
				"key": "agreements",
				"title": "Договорённости",
				"description": "Что команда решила",
				"lesson": self.урок,
				"rows": "many",
				"columns": ЖУРНАЛ["columns"],
				"items": [
					{"id": "A1", "topic": "Ретро", "decision": "Раз в две недели", "owner": "Маша"},
					{"id": "A2", "topic": "Созвоны"},
				],
				"filled": True,
				"empty_cells": [],
			},
		)
		self.assertEqual(
			(правила["rows"], правила["values"], правила["filled"]), ("one", {"method": "Согласие"}, True)
		)
		self.assertNotIn("items", правила)
		self.assertNotIn("values", журнал)

		with self.subTest("перечень и прогресс"):
			[перечень] = student.artifact(self.курс)["data"]["artifacts"]
			self.assertEqual(
				{к: перечень[к] for к in ("artifact", "purpose", "blocks_total", "blocks_filled")},
				{
					"artifact": "agreements_log",
					"purpose": "Зачем журнал",
					"blocks_total": 2,
					"blocks_filled": 2,
				},
			)
			курс = next(к for к in student.get_my_progress()["data"]["courses"] if к["id"] == self.курс)
			self.assertEqual(курс["documents"], [перечень])

	def test_очистка(self):
		"""Очистка незаведённого документа — не отказ и документа не заводит;
		очистка раздела стирает только его."""
		ответ = self.записать("agreements", clear="true")
		self.assertTrue(ответ["ok"], ответ)
		self.assertFalse(
			frappe.db.exists("Agent Student Artifact", {"student": self.ученик, "course": self.курс})
		)

		self.записать("agreements", rows=[{"topic": "Ретро"}])
		self.записать("how_we_decide", fields={"method": "Согласие"})
		ответ = self.записать("agreements", clear=True)["data"]
		self.assertEqual((ответ["blocks_filled"], ответ["created"]), (1, []))
		журнал, правила = self.документ()["blocks"]
		self.assertEqual((журнал["items"], правила["values"]), ([], {"method": "Согласие"}))
		self.assertEqual(self.записать("agreements", rows=[{"topic": "Новая"}])["data"]["created"], ["A2"])

	def test_отказ_записи_доходит_до_метода(self):
		"""Проводка: отказ сервиса — кодом контракта. Отказы разбора строк и полей
		— в `test_tables`, отказы самой записи — `test_отказы_записи`."""
		ответ = self.записать("agreements", fields={"topic": "x"})
		self.assertEqual(ответ["error"]["code"], codes.НЕ_ТОТ_РАЗДЕЛ)

	def test_только_дописать_доходит_до_записи(self):
		"""Проводка `fill_only` (#614): дописывание проходит, правка заполненного —
		отказ со списком затронутого. Отказы самой записи — `test_отказы_записи`."""
		self.assertTrue(self.записать("agreements", rows=[{"topic": "Ретро"}], fill_only=True)["ok"])

		ответ = self.записать("agreements", rows=[{"id": "A1", "topic": "Демо"}], fill_only=True)

		self.assertEqual(
			(ответ["error"]["code"], ответ["error"]["touched"]), (codes.ПРАВКА_ЗАПОЛНЕННОГО, ["A1.topic"])
		)

	def test_отказы_записи(self):
		"""Отказы, которые даёт сама запись (`write.записать_блок`), — до записи."""
		self.записать("agreements", rows=[{"topic": "Ретро"}])
		for что, (параметры, код) in {
			"нет документа": ({"artifact": "lean_canvas", "key": "agreements"}, codes.АРТЕФАКТ_НЕ_НАЙДЕН),
			"нет блока": ({"key": "budget", "rows": [{"topic": "x"}]}, codes.БЛОК_НЕ_НАЙДЕН),
			"очистка с записью": (
				{"key": "agreements", "rows": [{"topic": "x"}], "clear": True},
				codes.ОЧИСТКА_С_ЗАПИСЬЮ,
			),
			# Только дописать (#614): заполненное, удаление и очистка — отказ.
			"дописать: правка": (
				{"key": "agreements", "rows": [{"id": "A1", "topic": "Демо"}], "fill_only": True},
				codes.ПРАВКА_ЗАПОЛНЕННОГО,
			),
			"дописать: удаление": (
				{"key": "agreements", "delete_rows": ["A1"], "fill_only": True},
				codes.ПРАВКА_ЗАПОЛНЕННОГО,
			),
			"дописать: очистка": (
				{"key": "agreements", "clear": True, "fill_only": True},
				codes.ПРАВКА_ЗАПОЛНЕННОГО,
			),
		}.items():
			with self.subTest(что):
				with self.assertRaises(Отказ) as отказ:
					write.записать_блок(
						self.ученик, self.курс, None, **{"artifact": "agreements_log", **параметры}
					)
				self.assertEqual(отказ.exception.код, код)
		self.assertEqual(self.документ()["blocks"][0]["items"], [{"id": "A1", "topic": "Ретро"}])

	def test_правка_схемы_не_теряет_данных_и_помнит_версию(self):
		self.записать("how_we_decide", fields={"method": "Согласие"})
		вторая = self.схема([ЖУРНАЛ])
		self.assertEqual([б["key"] for б in self.документ()["blocks"]], ["agreements"])
		self.записать("agreements", rows=[{"topic": "Ретро"}])
		self.assertEqual(
			frappe.db.get_value(
				"Agent Student Artifact", {"student": self.ученик, "course": self.курс}, "schema_version"
			),
			вторая,
		)
		self.схема([ЖУРНАЛ, ПРАВИЛА])
		self.assertEqual(self.документ()["blocks"][1]["values"], {"method": "Согласие"})

		with self.subTest("документы — в порядке объявления"):
			self.схема([ПРАВИЛА], ключ="map")
			self.схема([ЖУРНАЛ, ПРАВИЛА])
			self.assertEqual(
				[а["artifact"] for а in student.artifact(self.курс)["data"]["artifacts"]],
				["agreements_log", "map"],
			)

	def test_запись_двигает_время_и_версию_и_шлёт_событие(self):
		"""`modified` — `null`, пока документ пуст, затем время записи со смещением
		пояса (#462); сколько раз документ сохранялся — без самой истории (#342);
		событие для открытой страницы — только ученику и после коммита (#348)."""
		пустой = self.документ()
		self.assertEqual((пустой["modified"], пустой["version"]), (None, 0))
		self.assertIsNone(student.artifact(self.курс)["data"]["artifacts"][0]["modified"])

		with patch("frappe.publish_realtime") as publish:
			self.записать("agreements", rows=[{"topic": "Ретро"}])

		события = [c for c in publish.call_args_list if c.kwargs.get("event") == "artifact_updated"]
		self.assertEqual(len(события), 1)
		self.assertEqual(
			(
				события[0].kwargs["user"],
				события[0].kwargs["message"]["artifact"],
				события[0].kwargs["after_commit"],
			),
			(self.ученик, "agreements_log", True),
		)
		перечень = student.artifact(self.курс)["data"]["artifacts"][0]
		self.assertIsNotNone(datetime.fromisoformat(перечень["modified"]).tzinfo)
		документ = self.документ()
		self.assertEqual(документ["modified"], перечень["modified"])
		self.assertIsInstance(документ["version"], int)
		self.assertNotIn("versions", документ)

	def test_выгрузка_markdown_и_xlsx(self):
		"""Проводка выгрузки: имя файла и листы книги; содержимое — в `test_tables`."""
		from openpyxl import load_workbook

		from lms_frappe_app.www.artifacts import download

		download(self.курс, "agreements_log")
		self.assertEqual(frappe.response["filename"], "agreements_log.md")
		self.assertTrue(frappe.response["filecontent"].startswith("# Журнал договорённостей\n"))

		# Подписи книги — на языке запроса выгрузившего (learning-services#767).
		self.addCleanup(setattr, frappe.local, "lang", frappe.local.lang)
		for язык, лист in (("ru", "Поля"), ("en", "Fields")):
			with self.subTest(язык):
				frappe.local.lang = язык
				download(self.курс, "agreements_log", format="xlsx")
				self.assertEqual(frappe.response["filename"], "agreements_log.xlsx")
				книга = load_workbook(BytesIO(frappe.response["filecontent"]))
				self.assertEqual(книга.sheetnames, ["Договорённости", лист])
