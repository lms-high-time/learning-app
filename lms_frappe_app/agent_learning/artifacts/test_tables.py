# Copyright (c) 2026, NikoMusaev and contributors
# For license information, please see license.txt

"""Разделы документа без базы: схема, запись, заполненность, выгрузка
(learning-services#515).

Схема — журнал договорённостей в миниатюре: раздел `many` с условной
колонкой («кто отвечает» — когда есть решение) и раздел `one` с полями.
"""

import io

from frappe.tests import UnitTestCase

from lms_frappe_app.agent_learning.artifacts import codes, data, export, fill, schema
from lms_frappe_app.agent_learning.errors import Отказ

ЖУРНАЛ = schema.проверить_раздел(
	{
		"key": "agreements",
		"title": "Договорённости",
		"rows": "many",
		"columns": [
			{"key": "topic", "title": "Тема", "required": True},
			{"key": "decision", "title": "Решение", "required": False},
			{"key": "owner", "title": "Кто отвечает", "required": {"if_column": "decision"}},
		],
	}
)
ПРАВИЛА = schema.проверить_раздел(
	{
		"key": "how_we_decide",
		"title": "Как мы решаем",
		"rows": "one",
		"columns": [
			{"key": "method", "title": "Способ", "required": True},
			{"key": "exception", "title": "Исключение", "required": False},
			{"key": "who_breaks", "title": "Кто снимает", "required": {"if_column": "exception"}},
		],
	}
)


def отказ(тест, код: str, действие, **подробности) -> None:
	with тест.assertRaises(Отказ) as пойман:
		действие()
	тест.assertEqual(пойман.exception.код, код, пойман.exception.сообщение)
	for ключ, значение in подробности.items():
		тест.assertEqual(пойман.exception.подробности.get(ключ), значение, ключ)


class TestСхема(UnitTestCase):
	def test_раздел_в_каноническом_виде(self):
		self.assertEqual(
			ЖУРНАЛ,
			{
				"key": "agreements",
				"title": "Договорённости",
				"description": None,
				"lesson": None,
				"rows": "many",
				"columns": [
					{"key": "topic", "title": "Тема", "required": True},
					{"key": "decision", "title": "Решение", "required": False},
					{"key": "owner", "title": "Кто отвечает", "required": {"if_column": "decision"}},
				],
			},
		)
		self.assertEqual(
			schema.проверить_раздел({"key": " Notes ", "rows": "one", "columns": [{"key": "text"}]}),
			{
				"key": "notes",
				"title": "notes",
				"description": None,
				"lesson": None,
				"rows": "one",
				"columns": [{"key": "text", "title": "text", "required": False}],
			},
		)

	def test_неверный_раздел_отказ(self):
		def раздел(**правка):
			return {"key": "log", "rows": "many", "columns": [{"key": "topic"}], **правка}

		for имя, описание in {
			"ключ не латиницей": раздел(key="Лог"),
			"вид строк не из двух": раздел(rows="table"),
			"колонки не список": раздел(columns={"topic": {}}),
			"колонка id": раздел(columns=[{"key": "id"}]),
			"ключ колонки с заглавной": раздел(columns=[{"key": "Topic"}]),
			"колонка дважды": раздел(columns=[{"key": "topic"}, {"key": "topic"}]),
			"условие на чужую колонку": раздел(columns=[{"key": "a", "required": {"if_column": "b"}}]),
			"условие на себя": раздел(columns=[{"key": "a", "required": {"if_column": "a"}}]),
			"required строкой": раздел(columns=[{"key": "a"}, {"key": "b", "required": "a"}]),
			"лишний ключ условия": раздел(
				columns=[{"key": "a"}, {"key": "b", "required": {"if_column": "a", "x": 1}}]
			),
		}.items():
			with self.subTest(имя):
				отказ(self, codes.НЕВЕРНАЯ_СХЕМА, lambda описание=описание: schema.проверить_раздел(описание))

	def test_документ_целиком(self):
		def раздел(ключ: str, rows: str, *колонки: str) -> dict:
			return schema.проверить_раздел(
				{"key": ключ, "rows": rows, "columns": [{"key": к} for к in колонки]}
			)

		schema.проверить_документ(
			[раздел("a", "many", "topic"), раздел("b", "many", "topic"), раздел("c", "one", "x")]
		)
		for имя, разделы in {
			"раздел дважды": [раздел("a", "one", "x"), раздел("a", "many", "y")],
			"поле совпало с полем": [раздел("a", "one", "x"), раздел("b", "one", "x")],
			"поле совпало с колонкой": [раздел("a", "many", "x"), раздел("b", "one", "x")],
		}.items():
			with self.subTest(имя):
				отказ(self, codes.НЕВЕРНАЯ_СХЕМА, lambda разделы=разделы: schema.проверить_документ(разделы))


class TestЗапись(UnitTestCase):
	def test_строки_заводятся_правятся_и_удаляются(self):
		д = data.записать(
			ЖУРНАЛ, None, rows=[{"topic": "  Ретро раз в две недели  "}, {"topic": "Созвоны", "decision": 12}]
		)
		self.assertEqual(
			д["tables"]["agreements"],
			[
				{"id": "A1", "topic": "Ретро раз в две недели"},
				{"id": "A2", "topic": "Созвоны", "decision": "12"},
			],
		)
		д = data.записать(ЖУРНАЛ, д, rows=[{"id": "a2", "decision": "", "owner": "Маша"}])
		self.assertEqual(д["tables"]["agreements"][1], {"id": "A2", "topic": "Созвоны", "owner": "Маша"})
		д = data.записать(ЖУРНАЛ, д, delete_rows=["A2"], rows=[{"topic": "Третья"}])
		self.assertEqual(
			[с["id"] for с in д["tables"]["agreements"]], ["A1", "A3"], "номера не переиспользуются"
		)
		self.assertEqual(д["seq"], {"agreements": 3})

	def test_поля_пишутся_и_стираются(self):
		д = data.записать(ПРАВИЛА, None, fields={"method": " Согласие ", "exception": "Бюджет"})
		д = data.записать(ПРАВИЛА, д, fields={"exception": "  "})
		self.assertEqual(д["fields"], {"method": "Согласие"})

	def test_отказы(self):
		журнал = data.записать(ЖУРНАЛ, None, rows=[{"topic": "Ретро"}])
		длинно = "я" * 4001
		for имя, (раздел, параметры, код, подробности) in {
			"строки у раздела one": (
				ПРАВИЛА,
				{"rows": [{"method": "x"}]},
				codes.НЕ_ТОТ_РАЗДЕЛ,
				{"rows": "one"},
			),
			"удаление у раздела one": (
				ПРАВИЛА,
				{"delete_rows": ["H1"]},
				codes.НЕ_ТОТ_РАЗДЕЛ,
				{"rows": "one"},
			),
			"поля у раздела many": (
				ЖУРНАЛ,
				{"fields": {"topic": "x"}},
				codes.НЕ_ТОТ_РАЗДЕЛ,
				{"rows": "many"},
			),
			"нечего писать": (ЖУРНАЛ, {"rows": [], "fields": {}}, codes.НЕЧЕГО_ПИСАТЬ, {}),
			"нет колонки": (
				ЖУРНАЛ,
				{"rows": [{"topik": "x"}]},
				codes.НЕТ_КОЛОНКИ,
				{"column": "topik", "columns": ["topic", "decision", "owner"]},
			),
			"нет поля": (
				ПРАВИЛА,
				{"fields": {"metod": "x"}},
				codes.НЕТ_ПОЛЯ,
				{"field": "metod", "fields": ["method", "exception", "who_breaks"]},
			),
			"правка строки, которой нет": (
				ЖУРНАЛ,
				{"rows": [{"id": "A9", "topic": "x"}]},
				codes.НЕТ_СТРОКИ,
				{"row": "A9"},
			),
			"удаление строки, которой нет": (
				ЖУРНАЛ,
				{"delete_rows": ["A9"]},
				codes.НЕТ_СТРОКИ,
				{"row": "A9"},
			),
			"длинная клетка": (
				ЖУРНАЛ,
				{"rows": [{"topic": длинно}]},
				codes.НЕВЕРНОЕ_ЗНАЧЕНИЕ,
				{"column": "topic", "max": 4000},
			),
			"длинное поле": (
				ПРАВИЛА,
				{"fields": {"method": длинно}},
				codes.НЕВЕРНОЕ_ЗНАЧЕНИЕ,
				{"field": "method", "max": 4000},
			),
			"не текст": (ЖУРНАЛ, {"rows": [{"topic": ["a"]}]}, codes.НЕВЕРНОЕ_ЗНАЧЕНИЕ, {"column": "topic"}),
			"флажок не текст": (
				ПРАВИЛА,
				{"fields": {"method": True}},
				codes.НЕВЕРНОЕ_ЗНАЧЕНИЕ,
				{"field": "method"},
			),
			"строки не список": (ЖУРНАЛ, {"rows": "{}"}, codes.НЕВЕРНОЕ_ЗНАЧЕНИЕ, {"param": "rows"}),
		}.items():
			with self.subTest(имя):
				до = repr(журнал)
				отказ(
					self,
					код,
					lambda раздел=раздел, параметры=параметры: data.записать(раздел, журнал, **параметры),
					**подробности,
				)
				self.assertEqual(repr(журнал), до, "отказ не меняет данных")

	def test_больше_предела_строк_отказ(self):
		д = data.записать(ЖУРНАЛ, None, rows=[{"topic": str(н)} for н in range(data.СТРОК_МАКС)])
		отказ(
			self,
			codes.СТРОК_СЛИШКОМ_МНОГО,
			lambda: data.записать(ЖУРНАЛ, д, rows=[{"topic": "лишняя"}]),
			max=data.СТРОК_МАКС,
		)

	def test_строки_и_json_строкой(self):
		"""Тело запроса формой приносит списки и объекты строками JSON."""
		д = data.записать(ЖУРНАЛ, None, rows='[{"topic": "Ретро"}]')
		д = data.записать(ПРАВИЛА, д, fields='{"method": "Согласие"}')
		self.assertEqual(
			(д["tables"]["agreements"][0]["topic"], д["fields"]["method"]), ("Ретро", "Согласие")
		)

	def test_очистка(self):
		д = data.записать(ЖУРНАЛ, None, rows=[{"topic": "Ретро"}])
		д = data.записать(ПРАВИЛА, д, fields={"method": "Согласие"})
		д["fields"]["other_field"] = "чужое"

		без_журнала = data.очистить(ЖУРНАЛ, д)
		self.assertNotIn("agreements", без_журнала["tables"])
		self.assertEqual(без_журнала["seq"], {"agreements": 1}, "номера после очистки не переиспользуются")
		self.assertEqual(data.очистить(ПРАВИЛА, д)["fields"], {"other_field": "чужое"})


class TestЗаполненность(UnitTestCase):
	def test_раздел_many(self):
		for имя, строки, заполнен, пустые in (
			("строк нет", [], False, []),
			("обязательное непусто", [{"topic": "Ретро"}], True, []),
			(
				"нет обязательного и условного",
				[{"decision": "Раз в две недели"}],
				False,
				[{"row": "A1", "column": "topic"}, {"row": "A1", "column": "owner"}],
			),
			("условие выполнено", [{"topic": "Ретро", "decision": "Да", "owner": "Маша"}], True, []),
		):
			with self.subTest(имя):
				д = data.записать(ЖУРНАЛ, None, rows=строки) if строки else data.данные(None)
				self.assertEqual(fill.заполнен(ЖУРНАЛ, д), заполнен)
				self.assertEqual(fill.пустые_клетки(ЖУРНАЛ, д), пустые)

	def test_раздел_one(self):
		нет_обязательных = schema.проверить_раздел(
			{"key": "notes", "rows": "one", "columns": [{"key": "note"}]}
		)
		for имя, раздел, поля, заполнен, пустые in (
			("пусто", ПРАВИЛА, {}, False, [{"field": "method"}]),
			("обязательное есть", ПРАВИЛА, {"method": "Согласие"}, True, []),
			(
				"условное не хватает",
				ПРАВИЛА,
				{"method": "Согласие", "exception": "Бюджет"},
				False,
				[{"field": "who_breaks"}],
			),
			("без обязательных — пусто", нет_обязательных, {}, False, []),
			("без обязательных — есть значение", нет_обязательных, {"note": "x"}, True, []),
		):
			with self.subTest(имя):
				д = data.данные({"fields": поля})
				self.assertEqual(fill.заполнен(раздел, д), заполнен)
				self.assertEqual(fill.пустые_клетки(раздел, д), пустые)

	def test_колонка_вне_схемы_не_видна(self):
		"""Клетка колонки, убранной из схемы, хранится, но в строках раздела её нет."""
		д = data.данные(
			{"tables": {"agreements": [{"id": "A1", "topic": "Ретро", "old": "x", "decision": ""}]}}
		)
		self.assertEqual(data.строки(ЖУРНАЛ, д), [{"id": "A1", "topic": "Ретро"}])


class TestВыгрузка(UnitTestCase):
	def данные(self) -> dict:
		д = data.записать(ЖУРНАЛ, None, rows=[{"topic": "Ретро | раз\nв две недели", "decision": "=1+1"}])
		return data.записать(ПРАВИЛА, д, fields={"method": "Согласие\nили голосование"})

	def test_markdown(self):
		д = self.данные()
		self.assertEqual(
			export.markdown(ЖУРНАЛ, д),
			"| ID | Тема | Решение | Кто отвечает |\n| --- | --- | --- | --- |\n"
			"| A1 | Ретро \\| раз в две недели | =1+1 |  |",
		)
		self.assertEqual(export.markdown(ПРАВИЛА, д), "- Способ: Согласие\n  или голосование")
		self.assertEqual(export.markdown(ЖУРНАЛ, data.данные(None)), "")
		self.assertEqual(
			export.собрать_markdown("Журнал", [ПРАВИЛА, ЖУРНАЛ], data.данные(None)),
			"# Журнал\n\n## Как мы решаем\n\n_Не заполнено._\n\n## Договорённости\n\n_Не заполнено._\n",
		)

	def test_книга_xlsx(self):
		from openpyxl import load_workbook

		второй = {**ЖУРНАЛ, "key": "areas", "title": "Договорённости"}
		странный = {**ЖУРНАЛ, "key": "odd", "title": "Что/где: [итог]?" + "я" * 40}
		книга = load_workbook(
			io.BytesIO(export.книга_xlsx([ЖУРНАЛ, ПРАВИЛА, второй, странный], self.данные()))
		)

		self.assertEqual(
			книга.sheetnames,
			["Договорённости", "Поля", "Договорённости (2)", "Что где   итог  " + "я" * 15],
		)
		журнал = книга["Договорённости"]
		self.assertEqual([к.value for к in журнал[1]], ["ID", "Тема", "Решение", "Кто отвечает"])
		self.assertEqual([к.value for к in журнал[2]][:3], ["A1", "Ретро | раз\nв две недели", "=1+1"])
		self.assertEqual(журнал["C2"].data_type, "s", "текст с «=» — не формула")
		self.assertEqual(
			[[к.value for к in ряд] for ряд in книга["Поля"].iter_rows()],
			[
				["Раздел", "Поле", "Значение"],
				["Как мы решаем", "Способ", "Согласие\nили голосование"],
				["Как мы решаем", "Исключение", None],
				["Как мы решаем", "Кто снимает", None],
			],
		)
