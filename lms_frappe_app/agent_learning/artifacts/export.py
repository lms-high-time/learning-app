# Copyright (c) 2026, NikoMusaev and contributors
# For license information, please see license.txt

"""Документ наружу: markdown раздела и документа, книга Excel."""

from __future__ import annotations

import re

from lms_frappe_app.agent_learning.artifacts.data import поля, строки
from lms_frappe_app.agent_learning.artifacts.schema import МНОГО

#: Чем markdown документа отмечает пустой раздел.
ПУСТОЙ_РАЗДЕЛ = "_Не заполнено._"
#: Лист книги с полями разделов `one`.
ЛИСТ_ПОЛЕЙ = "Поля"
#: Знаки, которых Excel не принимает в имени листа, и длина имени.
НЕЛЬЗЯ_В_ЛИСТЕ = re.compile(r"[\[\]:*?/\\]")
ДЛИНА_ЛИСТА = 31
#: Имя листа, которое Excel держит за собой: книгу с таким листом он не открывает.
ЗАНЯТОЕ_ИМЯ_ЛИСТА = "history"


def _текст(значение) -> str:
	"""Значение клетки или поля текстом. `Why:` движок пишет только текст, но
	данные, записанные раньше, могут хранить число или флажок."""
	return "" if значение is None else str(значение)


def markdown(раздел: dict, д: dict) -> str:
	"""Раздел текстом: у `many` — таблица с колонкой ID, у `one` — поля
	строками «Название: значение»; пустой раздел — пустая строка."""
	if раздел["rows"] == МНОГО:
		return _таблица(раздел["columns"], строки(раздел, д))
	значения = поля(раздел, д)
	return "\n".join(
		f"- {к['title']}: " + _текст(значения[к["key"]]).replace("\n", "\n  ")
		for к in раздел["columns"]
		if к["key"] in значения
	)


def _таблица(колонки: list[dict], ряды: list[dict]) -> str:
	if not ряды:
		return ""

	def клетка(значение) -> str:
		return _текст(значение).replace("|", "\\|").replace("\n", " ")

	шапка = ["ID", *(к["title"] for к in колонки)]
	линии = ["| " + " | ".join(клетка(ш) for ш in шапка) + " |", "|" + " --- |" * len(шапка)]
	for р in ряды:
		линии.append(
			"| " + " | ".join(клетка(р.get(к, "")) for к in ["id", *(к["key"] for к in колонки)]) + " |"
		)
	return "\n".join(линии)


def собрать_markdown(название: str, разделы: list[dict], д: dict) -> str:
	"""Документ одним файлом: заголовок и разделы в порядке схемы."""
	части = [f"# {название}"]
	for раздел in разделы:
		части.append(f"## {раздел['title']}")
		части.append(markdown(раздел, д) or ПУСТОЙ_РАЗДЕЛ)
	return "\n\n".join(части) + "\n"


def книга_xlsx(разделы: list[dict], д: dict) -> bytes:
	"""Документ книгой Excel: лист на раздел `many` и лист «Поля» с полями
	разделов `one` — на месте первого из них. Пустые клетки и поля — тоже в
	книге: её продолжают вести после курса.

	Текст, начатый с `=`, — и в клетке, и в шапке — остаётся текстом. `Why:`
	иначе Excel счёл бы его формулой и посчитал при открытии — текст ученика
	или автора исполнялся бы у того, кому передали книгу.
	"""
	import io

	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font, PatternFill
	from openpyxl.utils import get_column_letter

	книга = Workbook()
	книга.remove(книга.active)
	заливка = PatternFill("solid", fgColor="EEF2F6")
	занятые: set[str] = set()
	лист_полей = None

	def текстом(клетки) -> None:
		for клетка in клетки:
			if isinstance(клетка.value, str) and клетка.value.startswith("="):
				клетка.data_type = "s"

	def шапка(лист, названия: list[str], ширины: list[int]) -> None:
		лист.append(названия)
		текстом(лист[1])
		for клетка in лист[1]:
			клетка.font = Font(bold=True)
			клетка.fill = заливка
			клетка.alignment = Alignment(wrap_text=True, vertical="top")
		for номер, ширина in enumerate(ширины, start=1):
			лист.column_dimensions[get_column_letter(номер)].width = ширина
		лист.freeze_panes = "A2"

	def строка(лист, значения: list) -> None:
		лист.append([None if значение is None else _текст(значение) for значение in значения])
		текстом(лист[лист.max_row])
		for клетка in лист[лист.max_row]:
			клетка.alignment = Alignment(wrap_text=True, vertical="top")

	for раздел in разделы:
		if раздел["rows"] == МНОГО:
			лист = книга.create_sheet(_имя_листа(раздел["title"], занятые))
			шапка(
				лист, ["ID", *(к["title"] for к in раздел["columns"])], [8, *(36 for _ in раздел["columns"])]
			)
			for р in строки(раздел, д):
				строка(лист, [р["id"], *(р.get(к["key"]) for к in раздел["columns"])])
			continue
		if лист_полей is None:
			лист_полей = книга.create_sheet(_имя_листа(ЛИСТ_ПОЛЕЙ, занятые))
			шапка(лист_полей, ["Раздел", "Поле", "Значение"], [30, 36, 60])
		значения = поля(раздел, д)
		for к in раздел["columns"]:
			строка(лист_полей, [раздел["title"], к["title"], значения.get(к["key"])])
	if not книга.worksheets:
		книга.create_sheet(ЛИСТ_ПОЛЕЙ)

	поток = io.BytesIO()
	книга.save(поток)
	return поток.getvalue()


def _имя_листа(название: str, занятые: set[str]) -> str:
	"""Имя листа Excel из названия раздела: без запрещённых знаков, не длиннее
	31 знака, не повторяется в книге (без учёта регистра, как сравнивает Excel)
	и не совпадает с именем, которое Excel держит за собой."""
	основа = НЕЛЬЗЯ_В_ЛИСТЕ.sub(" ", название or "").strip().strip("'") or "Раздел"
	имя = основа[:ДЛИНА_ЛИСТА]
	номер = 1
	while имя.lower() in занятые or имя.lower() == ЗАНЯТОЕ_ИМЯ_ЛИСТА:
		номер += 1
		хвост = f" ({номер})"
		имя = основа[: ДЛИНА_ЛИСТА - len(хвост)] + хвост
	занятые.add(имя.lower())
	return имя
